#!/usr/bin/env python3
"""
Build the NDIS VA Foundations static site from the RC4 learner course files.

Content source of truth: Learner_Course/*.md (RC4). This script does not
rewrite course content. It only:
  * splits each module into preamble / lessons (steps) / wrap-up sections,
  * renders Markdown to HTML,
  * strips non-content authoring artifacts (chat-tool citation markers),
  * normalises Pandoc-style "---"/"--" dashes to typographic dashes,
  * adds interaction around existing activities (fillable empty table cells,
    checklists, response boxes on Try It / Check Yourself / Prove It),
  * excludes one out-of-place block (Module 11 intro text inside Module 2),
which is logged and reported.
"""
import html
import json
import os
import re
import sys
from pathlib import Path

import markdown
import openpyxl

SRC = Path(sys.argv[1])            # .../NDIS_VA_Foundations_RC4
OUT = Path(sys.argv[2])            # site root
PACK = SRC / "Wattlebird" / "Wattlebird_Sandbox_Pack_v2.3.4.xlsx"
PACK_FILE = "downloads/Wattlebird_Sandbox_Pack_v2.3.4.xlsx"
# Authoritative learner workbook for M15 capstone execution (added to build inputs 27 Sep 2026).
M15_WB = None  # set in main() from argv[3]
M15_FILE = "downloads/M15_Learner_Synchronized_Rebuilt.xlsx"

# One place to switch on the certification route (see Build Report).
SUBMISSION_URL = ""

LOG = {"citations_removed": {}, "excluded_blocks": [], "lesson_intro": [], "warnings": []}

MODULE_TITLES = {
    1: "Understanding the NDIS VA Work Environment",
    2: "Workplace Communication",
    3: "Tools, Organization, and Basic Prioritization",
    4: "Time Zones, Australian Formats, and Reliability",
    5: "What the NDIS Is",
    6: "Plans, Support Budgets, and Plan Management",
    7: "Roles and Terminology",
    8: "Pricing Documents: Where to Check and What to Do",
    9: "Scheduling and Rosters",
    10: "Onboarding and Service Agreements",
    11: "Invoicing and Claims",
    12: "Privacy and Confidentiality",
    13: "Compliance Administration",
    14: "Systems, SOPs & Administrative Workflows",
    15: "Capstone Project — The Wattlebird Mixed Workweek",
    16: "Portfolio and Finding Clients",
    17: "Interviews and Your First 30 Days",
}

STEP_RE = re.compile(
    r"^([1-9])\.\s+(Progress Reminder|Why This Matters|Learn|See It.*|Try It.*|Check Yourself.*|Prove It.*|Work Boundary.*|Next Step.*)$"
)
STEP_SHORT = {
    "Progress Reminder": "Progress Reminder", "Why This Matters": "Why This Matters", "Learn": "Learn",
    "See It": "See It", "Try It": "Try It", "Check Yourself": "Check Yourself", "Prove It": "Prove It",
    "Work Boundary": "Work Boundary", "Next Step": "Next Step",
}
POST_RE = re.compile(
    r"^(Module|Official|Quick Reference|M\d+ Quick|Certification|Capstone Assessment|Final|Course Completion|"
    r"Portfolio Artifact|Important assessment|Current-source|Current official)",
    re.I,
)
HEAD_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
CITE_RE = re.compile(r"\s*cite(?:[^]*)*")


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def esc(s):
    return html.escape(s, quote=True)


def clean_source(text, mod):
    n = len(CITE_RE.findall(text))
    if n:
        LOG["citations_removed"][mod] = n
    text = CITE_RE.sub("", text)
    # Pandoc-style dashes (Module 4): " --- " → em dash, "am--12" → en dash
    text = re.sub(r"(?<=\S) --- (?=\S)", " — ", text)
    text = re.sub(r"(?<=[\w)])--(?=[\w(])", "–", text)
    return text


def inline_md(s):
    """Headings/titles: escape, then allow `code` and **bold**."""
    s = esc(s)
    s = re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", s)
    return s


def plain(s):
    return re.sub(r"[`*]", "", s)


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


# ---------------------------------------------------------------------------
# Markdown rendering with post-processing
# ---------------------------------------------------------------------------
class Ctx:
    """Per-page counters so every saved field gets a stable, unique key."""

    def __init__(self, page_id):
        self.page_id = page_id
        self.table = 0
        self.check = 0
        self.pre = 0


def relevel(md_text, base_level, target):
    """Shift heading levels inside a chunk so the shallowest becomes `target`."""
    lines = md_text.split("\n")
    levels = [len(m.group(1)) for l in lines if (m := HEAD_RE.match(l))]
    if not levels:
        return md_text
    mn = min(levels) if base_level is None else base_level
    out = []
    in_fence = False
    for l in lines:
        if l.strip().startswith("```"):
            in_fence = not in_fence
        m = HEAD_RE.match(l) if not in_fence else None
        if m:
            lvl = len(m.group(1))
            new = min(max(target + (lvl - mn), target), 5)
            out.append("#" * new + " " + m.group(2))
        else:
            out.append(l)
    return "\n".join(out)


URL_RE = re.compile(r"(?<![(<\[\"'=])\b(https?://[^\s<>|)\]]+)")


def autolink(md_text):
    out = []
    in_fence = False
    for l in md_text.split("\n"):
        if l.strip().startswith("```"):
            in_fence = not in_fence
            out.append(l)
            continue
        if in_fence:
            out.append(l)
            continue

        def rep(m):
            url = m.group(1)
            trail = ""
            while url and url[-1] in ".,;:":
                trail = url[-1] + trail
                url = url[:-1]
            return "<" + url + ">" + trail

        # don't touch URLs already inside markdown links: [x](url)
        parts = re.split(r"(\[[^\]]*\]\([^)]*\))", l)
        parts = [p if p.startswith("[") and "](" in p else URL_RE.sub(rep, p) for p in parts]
        out.append("".join(parts))
    return "\n".join(out)


MD_EXT = ["tables", "fenced_code", "sane_lists"]


def md_to_html(md_text, ctx):
    # 1) Pull <details> blocks out; render their bodies separately.
    details = []

    def det(m):
        summary = m.group(1).strip()
        inner = m.group(2)
        details.append((summary, inner))
        return "\n\nDETAILSPLACEHOLDER%d\n\n" % (len(details) - 1)

    md_text = re.sub(r"<details>\s*<summary>(.*?)</summary>(.*?)</details>", det, md_text, flags=re.S)
    md_text = autolink(md_text)
    h = markdown.markdown(md_text, extensions=MD_EXT, output_format="html")

    def put_details(m):
        i = int(m.group(1))
        summary, inner = details[i]
        body = md_to_html(inner, ctx)
        return ('<details class="model-answer"><summary>%s</summary><div class="details-body md">%s</div></details>'
                % (inline_md(plain(summary) if "<" not in summary else summary), body))

    h = re.sub(r"<p>DETAILSPLACEHOLDER(\d+)</p>", put_details, h)
    return post_process(h, ctx)


def post_process(h, ctx):
    # External links open in a new tab and are marked as external.
    h = re.sub(r'<a href="(https?://[^"]+)">', r'<a class="ext-link" href="\1" target="_blank" rel="noopener">', h)

    # Blockquotes: fictional / important notices get the callout treatment.
    def bq(m):
        inner = m.group(1)
        if re.search(r"FICTIONAL|Important:|Practice rule|fictional training", inner, re.I):
            return '<blockquote class="callout-fictional">' + inner + "</blockquote>"
        return "<blockquote>" + inner + "</blockquote>"

    h = re.sub(r"<blockquote>(.*?)</blockquote>", bq, h, flags=re.S)

    # Current-information notes get a visible "check current source" frame.
    def cs(m):
        return ('<div class="current-source-note"><span class="cs-tag">Check the current official source</span>'
                + m.group(0) + "</div>")

    h = re.sub(r"<p><strong>(?:Current-information note|Important current-information note|Current-source reminder|"
               r"Current-information rule|Current information note)[^<]*</strong>.*?</p>", cs, h, flags=re.S)

    # Task-list checkboxes: "- [ ] text"
    def task_ul(m):
        ul = m.group(0)
        if not re.search(r"<li>\s*(?:<p>)?\[[ xX]\]", ul):
            return ul

        def li(mm):
            pre, text = mm.group(1), mm.group(2)
            ctx.check += 1
            key = "check-%d" % ctx.check
            label_text = re.sub(r"<[^>]+>", "", text).strip()
            return ('<li>%s<label><input type="checkbox" data-draft="%s" data-export-label="%s"> <span>%s</span></label>'
                    % (pre, key, esc(label_text), text))

        ul = re.sub(r"<li>(\s*(?:<p>)?)\[[ xX]\]\s*(.*?)(?=</p>|</li>)", li, ul, flags=re.S)
        return ul.replace("<ul>", '<ul class="task-list">', 1)

    h = re.sub(r"<ul>.*?</ul>", task_ul, h, flags=re.S)

    # Tables: wrap for scrolling, add labels; empty cells become inputs.
    def table(m):
        t = m.group(0)
        ctx.table += 1
        tn = ctx.table
        headers = [re.sub(r"<[^>]+>", "", x).strip() for x in re.findall(r"<th[^>]*>(.*?)</th>", t, flags=re.S)]
        rows = re.findall(r"<tr>(.*?)</tr>", t, flags=re.S)
        fillable = any(re.search(r"<td[^>]*>\s*</td>", r) for r in rows)
        out_rows = []
        body_row = 0
        for r in rows:
            if "<th" in r:
                out_rows.append("<tr>" + r + "</tr>")
                continue
            body_row += 1
            cells = re.findall(r"(<td[^>]*>)(.*?)</td>", r, flags=re.S)
            row_label = re.sub(r"<[^>]+>", "", cells[0][1]).strip() if cells else ""
            new_cells = []
            for ci, (open_tag, content) in enumerate(cells):
                col = headers[ci] if ci < len(headers) else "Column %d" % (ci + 1)
                attrs = ' data-label="%s"' % esc(col)
                align = re.search(r'style="[^"]*"', open_tag)
                if align:
                    attrs += " " + align.group(0)
                if content.strip() == "":
                    key = "t%d-r%d-c%d" % (tn, body_row, ci + 1)
                    lab = col + (" — " + row_label if row_label and ci > 0 else " — row %d" % body_row)
                    new_cells.append('<td class="fill-cell"%s><textarea rows="2" data-draft="%s" aria-label="%s" data-export-label="Table %d · %s"></textarea></td>'
                                     % (attrs, key, esc(lab), tn, esc(lab)))
                else:
                    new_cells.append("<td%s>%s</td>" % (attrs, content))
            out_rows.append("<tr>" + "".join(new_cells) + "</tr>")
        thead_rows = [r for r in out_rows if "<th" in r]
        tbody_rows = [r for r in out_rows if "<th" not in r]
        cls = "ref-table fill-table" if fillable else "ref-table"
        label = "Table: " + ", ".join(headers[:4]) if headers else "Table"
        inner = ("<thead>" + "".join(thead_rows) + "</thead>" if thead_rows else "") + "<tbody>" + "".join(tbody_rows) + "</tbody>"
        hint = '<p class="fill-hint">You can type in the empty cells — entries are saved in this browser only.</p>' if fillable else ""
        return ('<div class="table-wrap%s" role="region" aria-label="%s" tabindex="0"><table class="%s">%s</table></div>%s'
                % (" fill-wrap" if fillable else "", esc(label), cls, inner, hint))

    h = re.sub(r"<table>.*?</table>", table, h, flags=re.S)

    # Code-block templates get an id (for "copy into my notes").
    def pre(m):
        ctx.pre += 1
        return '<pre id="tpl-%s-%d" class="template-block">' % (ctx.page_id, ctx.pre)

    h = re.sub(r"<pre>", pre, h)
    return h


# ---------------------------------------------------------------------------
# Module parsing
# ---------------------------------------------------------------------------
def headings(lines):
    out = []
    in_fence = False
    for i, l in enumerate(lines):
        if l.strip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        m = HEAD_RE.match(l)
        if m:
            out.append((i, len(m.group(1)), m.group(2).strip()))
    return out


def strip_rules(text):
    """Remove leading/trailing horizontal rules and blank lines of a chunk."""
    lines = text.split("\n")
    while lines and (not lines[0].strip() or lines[0].strip() in ("---", "***")):
        lines.pop(0)
    while lines and (not lines[-1].strip() or lines[-1].strip() in ("---", "***")):
        lines.pop()
    return "\n".join(lines)


def parse_module(num, text):
    lines = text.split("\n")
    hs = headings(lines)
    lesson_re = re.compile(r"^(?:Lesson\s+)?%d\.(\d+)\s*[—–-]+\s*(.+)$" % num)

    # Title headings = first two headings.
    title_idx = {hs[0][0], hs[1][0]}

    lessons = []  # (start_line, level, lnum, title)
    for (i, lvl, t) in hs:
        m = lesson_re.match(t)
        if m and lvl <= 2:
            lessons.append((i, lvl, int(m.group(1)), m.group(2).strip()))
    if not lessons:
        raise SystemExit("No lessons in module %d" % num)

    first_lesson_line = lessons[0][0]
    lesson_level = lessons[0][1]
    lesson_lines = {l[0] for l in lessons}

    # Postamble start
    post_start = None
    for (i, lvl, t) in hs:
        if i <= first_lesson_line or i in lesson_lines:
            continue
        if STEP_RE.match(t):
            continue
        if lvl < lesson_level or (lvl == lesson_level) or (lvl <= 2 and POST_RE.match(t)):
            if lvl <= lesson_level or POST_RE.match(t):
                post_start = i
                break
    end_lessons = post_start if post_start is not None else len(lines)
    if any(l[0] > end_lessons for l in lessons):
        LOG["warnings"].append("Module %d: a lesson heading appears after wrap-up content." % num)

    pre_lines = [l for j, l in enumerate(lines[:first_lesson_line]) if j not in title_idx]
    preamble = strip_rules("\n".join(pre_lines))

    parsed_lessons = []
    for k, (start, lvl, lnum, title) in enumerate(lessons):
        stop = lessons[k + 1][0] if k + 1 < len(lessons) else end_lessons
        body_lines = lines[start + 1:stop]
        parsed_lessons.append(parse_lesson(num, lnum, title, body_lines))

    post = strip_rules("\n".join(lines[end_lessons:])) if post_start is not None else ""
    return preamble, parsed_lessons, post


def parse_lesson(mnum, lnum, title, body_lines):
    hs = headings(body_lines)
    step_level = None
    for (i, lvl, t) in hs:
        if re.match(r"^1\.\s+Progress Reminder", t):
            step_level = lvl
            break
    steps = []
    if step_level is not None:
        marks = [(i, t) for (i, lvl, t) in hs if lvl == step_level and STEP_RE.match(t)]
        intro = strip_rules("\n".join(body_lines[:marks[0][0]]))
        if intro.strip():
            LOG["lesson_intro"].append("%d.%d" % (mnum, lnum))
            steps.append({"heading": "Overview", "label": "Overview", "num": 0, "md": intro, "level": step_level})
        for k, (i, t) in enumerate(marks):
            stop = marks[k + 1][0] if k + 1 < len(marks) else len(body_lines)
            m = STEP_RE.match(t)
            label = STEP_SHORT[[s for s in STEP_SHORT if m.group(2).startswith(s)][0]]
            steps.append({"heading": t, "label": label, "num": int(m.group(1)),
                          "md": strip_rules("\n".join(body_lines[i + 1:stop])), "level": step_level})
    else:
        LOG["warnings"].append("Lesson %d.%d has no standard step headings; rendered as one panel." % (mnum, lnum))
        steps.append({"heading": "Lesson", "label": "Lesson", "num": 0, "md": strip_rules("\n".join(body_lines)), "level": None})
    return {"mnum": mnum, "lnum": lnum, "title": title, "steps": steps}


def split_sections(md_text):
    """Split a chunk at its shallowest headings → [(heading or None, body)]."""
    lines = md_text.split("\n")
    hs = headings(lines)
    if not hs:
        return [(None, md_text)] if md_text.strip() else []
    mn = min(l for (_, l, _) in hs)
    tops = [(i, t) for (i, l, t) in hs if l == mn]
    out = []
    head = strip_rules("\n".join(lines[:tops[0][0]]))
    if head.strip():
        out.append((None, head))
    for k, (i, t) in enumerate(tops):
        stop = tops[k + 1][0] if k + 1 < len(tops) else len(lines)
        out.append((t, strip_rules("\n".join(lines[i + 1:stop]))))
    return out


# ---------------------------------------------------------------------------
# Page chrome
# ---------------------------------------------------------------------------
HOME_SVG = '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M12 3l9 8h-3v9h-5v-6H11v6H6v-9H3l9-8z"/></svg>'


def page(root, title, description, crumbs, header_html, main_html, body_attrs="", scope=None):
    crumb_html = ""
    for c in crumbs:
        crumb_html += '<span class="sep" aria-hidden="true">/</span>'
        if c.get("href"):
            crumb_html += '<a class="crumb-link" href="%s">%s</a>' % (c["href"], esc(c["text"]))
        else:
            crumb_html += '<span class="current" aria-current="page">%s</span>' % esc(c["text"])
    scope_attr = ' data-draft-scope="%s"' % scope if scope else ""
    return """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>%(title)s</title>
<meta name="description" content="%(desc)s">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Fraunces:wght@500;600;700&family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="%(root)sshared/styles.css">
</head>
<body data-root="%(root)s"%(battrs)s>
<a class="skip-link" href="#main">Skip to content</a>

<nav class="route-bar" aria-label="Breadcrumb">
  <div class="route-track">
    <span class="crumb">PVA Free Training</span>
    <span class="sep" aria-hidden="true">/</span>
    <a class="home-btn" href="%(root)s">%(home)s NDIS VA Foundations</a>
    %(crumbs)s
  </div>
</nav>

%(header)s

<main id="main" class="container"%(scope)s>
%(main)s
</main>

<footer class="site-footer">
  <div class="container">
    <p>Probinsiyanong VA Free Training — genuinely free, no account required.</p>
    <p><a href="%(root)sprogress/">Your progress &amp; backup</a> · <a href="%(root)swattlebird/">Wattlebird practice file</a> · <a href="%(root)scertification-submission/">Certification Submission</a></p>
    <p>Wattlebird Support Services and all of its people, records and scenarios are fictional training data.</p>
  </div>
</footer>

<script src="%(root)sshared/course-data.js"></script>
<script src="%(root)sshared/progress-ndis.js"></script>
<script src="%(root)sshared/ui.js"></script>
</body>
</html>
""" % {"title": esc(title), "desc": esc(description), "root": root, "battrs": body_attrs, "home": HOME_SVG,
       "crumbs": crumb_html, "header": header_html, "scope": scope_attr, "main": main_html}


def header(eyebrow, h1, lede="", meta=""):
    return """<header class="page-header">
  <div class="container">
    <div class="eyebrow">%s</div>
    <h1>%s</h1>
    %s
    %s
  </div>
</header>""" % (eyebrow, h1, ('<p class="lede">%s</p>' % lede) if lede else "", ('<div class="header-meta">%s</div>' % meta) if meta else "")


def practice_strip(root, mnum=None):
    if mnum == 15:
        return ('<aside class="practice-strip no-print" aria-label="Capstone workbook">'
                '<span><strong>Capstone workbook:</strong> M15 Learner Workbook — contains <code>M15_Capstone</code> and the '
                'M15 capstone input sheets. Fictional training data. Work from a copy.</span>'
                '<span class="btn-row" style="margin:0;"><a class="btn btn-primary" href="%s%s" download>Download M15 workbook (.xlsx)</a>'
                '<a class="btn btn-secondary" href="%swattlebird/#m15-workbook">About the practice files</a></span></aside>' % (root, M15_FILE, root))
    return ('<aside class="practice-strip no-print" aria-label="Practice file">'
            '<span><strong>Practice file:</strong> Wattlebird Sandbox Pack v2.3.4 — fictional training data. Work from a copy.</span>'
            '<a class="btn btn-secondary" href="%swattlebird/">Open the Wattlebird practice file</a></aside>' % root)


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------
def mod_dir(n):
    return "module-%02d" % n


def lesson_dir(m, l):
    return "lesson-%d-%d" % (m, l)


def lesson_id(m, l):
    return "m%02d-l%d" % (m, l)


NOTES = {
    5: ("Your Try It work", "Work it here or in your own file. Saved in this browser only as you type."),
    6: ("Your answer — before you open the model answer", "Write your own answer first, then compare. Saved in this browser only."),
    7: ("Your Prove It notes", "Draft or track your Prove It evidence here. Saved in this browser only — save the finished evidence in your own file with the name the lesson gives you."),
}


def build_lesson(mods, mi, li, root="../../"):
    mod = mods[mi]
    les = mod["lessons"][li]
    m, l = mod["num"], les["lnum"]
    lid = lesson_id(m, l)
    ctx = Ctx(lid)
    label = "Lesson %d.%d" % (m, l)
    steps = les["steps"]
    panels = []
    n_panels = len(steps) + 1  # + Done
    notes_count = 0
    for si, st in enumerate(steps):
        md = st["md"]
        if st["level"] is not None:
            md = relevel(md, st["level"] + 1, 3)
        body = md_to_html(md, ctx)

        notes_html = ""
        if st["num"] in NOTES:
            notes_count += 1
            nl, help_text = NOTES[st["num"]]
            key = "notes-step%d" % st["num"]
            tid = "%s-%s" % (lid, key)
            notes_html = """<div class="notes-box no-print">
  <label for="%(tid)s">%(nl)s</label>
  <p class="notes-help">%(help)s</p>
  <textarea class="text-response" id="%(tid)s" data-draft="%(key)s" data-status="%(tid)s-status" data-export-label="%(nl)s"></textarea>
  <div class="save-status" id="%(tid)s-status" aria-live="polite"></div>
</div>""" % {"tid": tid, "nl": nl, "help": help_text, "key": key}
            # "Copy template into notes" for code-block templates in this panel
            for tpl in re.findall(r'<pre id="(tpl-[^"]+)"', body):
                body = body.replace('<pre id="%s"' % tpl,
                                    '<div class="btn-row no-print" style="margin:4px 0 8px;"><button type="button" class="btn btn-ghost" data-copy-template="%s" data-target="%s">Copy this template into my notes</button></div><pre id="%s"' % (tpl, tid, tpl), 1)
            if st["num"] == 6 and "<details" in body:
                pos = body.index("<details")
                body = body[:pos] + notes_html + body[pos:]
                notes_html = ""

        heading = inline_md(st["heading"]) if st["heading"] not in ("Overview", "Lesson") else st["heading"]
        nav_btns = []
        if si > 0:
            nav_btns.append('<button type="button" class="btn btn-ghost" data-goto="%d">← Back</button>' % (si - 1))
        else:
            nav_btns.append('<a class="btn btn-ghost" href="../">← Module %d overview</a>' % m)
        if si < len(steps) - 1:
            nav_btns.append('<button type="button" class="btn btn-primary" data-goto="%d">Continue: %s →</button>' % (si + 1, esc(steps[si + 1]["label"])))
        else:
            dl = "Mark %s as done" % label
            nav_btns.append('<button type="button" class="btn btn-primary" id="markDone" data-label="%s">%s</button>' % (esc(dl), esc(dl)))
        panels.append("""<section class="step-panel folder-tab" data-step="%d" data-label="%s" data-tab="%s" aria-labelledby="%s-h%d">
  <h2 id="%s-h%d">%s</h2>
  <div class="md">
%s
  </div>
  %s
  <div class="btn-row panel-nav no-print">%s</div>
</section>""" % (si, esc(st["label"]), esc(st["label"]), lid, si, lid, si, heading, body, notes_html, "".join(nav_btns)))

    # Next / previous lessons across modules
    flat = [(a, b) for a in range(len(mods)) for b in range(len(mods[a]["lessons"]))]
    pos = flat.index((mi, li))
    prev_link = next_link = ""
    if pos > 0:
        pa, pb = flat[pos - 1]
        pm, pl = mods[pa], mods[pa]["lessons"][pb]
        prev_link = '<a class="pager-prev" href="%s%s/%s/">← Lesson %d.%d — %s</a>' % (root, mod_dir(pm["num"]), lesson_dir(pm["num"], pl["lnum"]), pm["num"], pl["lnum"], esc(plain(pl["title"])))
    if pos < len(flat) - 1:
        na, nb = flat[pos + 1]
        nm, nl_ = mods[na], mods[na]["lessons"][nb]
        next_href = "%s%s/%s/" % (root, mod_dir(nm["num"]), lesson_dir(nm["num"], nl_["lnum"]))
        next_text = "Lesson %d.%d — %s" % (nm["num"], nl_["lnum"], plain(nl_["title"]))
        next_link = '<a class="pager-next" href="%s">%s →</a>' % (next_href, esc(next_text))
        cross = nm["num"] != m
        done_next = ('<a class="btn btn-primary" href="%s">%sStart %s →</a>' % (next_href, "Module %d · " % nm["num"] if cross else "", esc("Lesson %d.%d" % (nm["num"], nl_["lnum"]))))
    else:
        done_next = '<a class="btn btn-primary" href="%scertification-submission/">Go to Certification Submission →</a>' % root

    done_panel = """<section class="step-panel folder-tab" data-step="%(i)d" data-label="Done" data-tab="Done" data-state="done" aria-labelledby="%(lid)s-hdone">
  <span class="stamp"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M9 16.2l-3.5-3.6L4 14.1l5 5 11-11-1.5-1.4z"/></svg>%(label)s complete</span>
  <h2 id="%(lid)s-hdone" style="margin-top:14px;">%(label)s marked as done</h2>
  <p>Your completion is saved in this browser. Keep your finished Prove It work in your own files, too.</p>
  <div class="card" style="margin-top:16px;">
    <h3 style="font-size:1.05rem;">Keep a copy of your responses</h3>
    <p style="font-size:0.9rem;">Collect everything you typed in this lesson into one text you can copy, print or download.</p>
    <div class="btn-row" style="margin-top:0;"><button type="button" class="btn btn-secondary" id="generateNotes">Generate my %(label)s notes</button></div>
    <div class="export-box no-print" id="notesExportBox" hidden></div>
    <div class="btn-row no-print" id="notesExportActions" hidden><button type="button" class="btn btn-ghost" id="downloadNotes">Download as .txt</button></div>
    <p class="storage-note">Progress lives in this browser only. <a href="%(root)sprogress/">Export a progress backup</a> from time to time.</p>
  </div>
  <div class="lesson-nav-footer">
    <a class="btn btn-secondary" href="../">← Module %(m)d overview</a>
    %(next)s
  </div>
  <p style="margin-top:14px;"><button type="button" class="btn btn-ghost" id="markNotDone" hidden>Mark as not done</button></p>
</section>""" % {"i": n_panels - 1, "lid": lid, "label": label, "m": m, "next": done_next, "root": root}
    panels.append(done_panel)

    title_txt = plain(les["title"])
    eyebrow = "Module %d · Lesson %d of %d" % (m, li + 1, len(mod["lessons"]))
    if m == 15:
        eyebrow = '<span class="eyebrow-tag">Capstone assessment</span> &nbsp;' + eyebrow
    hdr = header(eyebrow, "%s — %s" % (label, inline_md(les["title"])), "", '<span data-lesson-state></span>')
    main = """%s
<nav class="step-nav no-print" id="stepNav" aria-label="Lesson steps"></nav>
<div class="progress-rail no-print" style="margin-top:6px;"><span id="lessonProgressBar" style="width:0%%;"></span></div>
%s
<nav class="lesson-pager no-print" aria-label="Lesson navigation">%s%s</nav>
""" % (practice_strip(root, m), "\n".join(panels), prev_link, next_link)
    attrs = ' data-lesson-id="%s" data-notes-title="%s" data-notes-filename="%s"' % (
        lid, esc("NDIS VA Foundations — %s — %s" % (label, title_txt)), "NDIS_M%d_Lesson_%d-%d_Notes.txt" % (m, m, l))
    crumbs = [{"text": "Module %d" % m, "href": "../"}, {"text": label}]
    return page(root, "%s: %s | NDIS VA Foundations — PVA Free Training" % (label, title_txt),
                "Module %d, %s of NDIS VA Foundations — %s." % (m, label, title_txt), crumbs, hdr, main, attrs, scope=lid)


def section_cards(md_text, ctx, cls="card", first_title=None, extra=None):
    out = []
    for (t, body) in split_sections(md_text):
        if extra and t and t in extra.get("skip", []):
            continue
        body_md = relevel(body, None, 3)
        body_html = md_to_html(body_md, ctx)
        title = inline_md(t) if t else (first_title or "")
        add = ""
        if extra and t and t in extra.get("after", {}):
            add = extra["after"][t]
        hid = "sec-" + slug(plain(t or first_title or "intro"))
        h2 = '<h2 id="%s">%s</h2>' % (hid, title) if title else ""
        out.append('<section class="%s" %s>%s<div class="md">%s</div>%s</section>'
                   % (cls, ('aria-labelledby="%s"' % hid) if title else "", h2, body_html, add))
    return "\n".join(out)


def build_module(mods, mi, root="../"):
    mod = mods[mi]
    m = mod["num"]
    ctx = Ctx(mod["id"])
    cert_btn = ('<div class="btn-row no-print"><a class="btn btn-primary" href="%scertification-submission/">'
                'Go to NDIS VA Foundations → Certification Submission</a></div>' % root)
    extra = {"skip": [], "after": {}}
    if m == 2:
        extra["skip"].append("Module 11 — Invoicing and Claims")
    if m == 15:
        extra["after"]["Certification Submission"] = cert_btn
    if m == 16:
        extra["after"]["Certification Submission Contract"] = cert_btn

    pre_html = section_cards(mod["preamble"], ctx, "card", first_title="About this module")
    lesson_cards = []
    for les in mod["lessons"]:
        lid = lesson_id(m, les["lnum"])
        lesson_cards.append("""<a class="lesson-card" href="%s/">
  <div class="lesson-num wide">%d.%d</div>
  <div class="lesson-info"><h3>%s</h3><p>%d steps · Learn → See It → Try It → Check → Prove It</p></div>
  <div class="lesson-status" data-lesson-status="%s"></div>
</a>""" % (lesson_dir(m, les["lnum"]), m, les["lnum"], inline_md(les["title"]), len([s for s in les["steps"] if s["num"]]), lid))

    post_html = ""
    if mod["post"]:
        post_html = ('<section class="folder-tab" data-tab="Module wrap-up" style="margin-top:36px;" aria-labelledby="wrap-h">'
                     '<h2 id="wrap-h">Module %d wrap-up</h2><p style="font-size:0.9rem;color:var(--ink-soft);">Practice, Prove It and reference material for the whole module. Checklist ticks are saved in this browser only.</p></section>\n' % m
                     + section_cards(mod["post"], ctx, "card", extra=extra))
    prev_m = ('<a class="pager-prev" href="%s%s/">← Module %d — %s</a>' % (root, mod_dir(mods[mi - 1]["num"]), mods[mi - 1]["num"], esc(mods[mi - 1]["title"]))) if mi > 0 else '<a class="pager-prev" href="%s">← Course home</a>' % root
    next_m = ('<a class="pager-next" href="%s%s/">Module %d — %s →</a>' % (root, mod_dir(mods[mi + 1]["num"]), mods[mi + 1]["num"], esc(mods[mi + 1]["title"]))) if mi + 1 < len(mods) else '<a class="pager-next" href="%scertification-submission/">Certification Submission →</a>' % root

    main = """%(strip)s
<section class="folder-tab" data-tab="Module %(m)d progress" data-state="current" aria-labelledby="mp-h">
  <h2 id="mp-h" style="font-size:1.1rem;">Your progress in this module</h2>
  <div class="progress-rail"><span data-module-bar="%(m)d" style="width:0%%;"></span></div>
  <div class="progress-meta"><span data-module-count="%(m)d"></span><span data-module-status="%(m)d"></span></div>
  <p class="storage-note">Saved in this browser only — see <a href="%(root)sprogress/">Your progress &amp; backup</a>.</p>
</section>
<div style="margin-top:24px;">%(pre)s</div>
<section style="margin-top:32px;" aria-labelledby="lessons-h">
  <h2 id="lessons-h">Lessons</h2>
  <div class="lesson-grid">%(lessons)s</div>
</section>
%(post)s
<nav class="lesson-pager no-print" aria-label="Module navigation">%(prev)s%(next)s</nav>
""" % {"strip": practice_strip(root, m), "m": m, "root": root, "pre": pre_html, "lessons": "\n".join(lesson_cards),
       "post": post_html, "prev": prev_m, "next": next_m}
    eyebrow = "Module %d of %d" % (m, len(mods))
    if m == 15:
        eyebrow = '<span class="eyebrow-tag">Capstone assessment</span> &nbsp;' + eyebrow
    hdr = header(eyebrow, "Module %d — %s" % (m, esc(mod["title"])), "", '<span data-module-status="%d"></span>' % m)
    crumbs = [{"text": "Module %d" % m}]
    return page(root, "Module %d: %s | NDIS VA Foundations — PVA Free Training" % (m, mod["title"]),
                "Module %d of NDIS VA Foundations — %s." % (m, mod["title"]), crumbs, hdr, main, scope=mod["id"])


# ---------------------------------------------------------------------------
# Workbook extracts (learner-facing sheets only)
# ---------------------------------------------------------------------------
def sheet_rows(ws):
    rows = []
    for r in ws.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() for v in r):
            rows.append(["" if v is None else (v.strftime("%-d %b %Y") if hasattr(v, "strftime") else str(v)) for v in r])
    return rows


def main():
    learner = SRC / "Learner_Course"
    mods = []
    for n in range(1, 18):
        f = sorted(learner.glob("NDIS_VA_Foundations_Module_%d_Learner_Facing_*.md" % n))
        assert len(f) == 1, f
        text = clean_source(f[0].read_text(encoding="utf-8"), n)
        preamble, lessons, post = parse_module(n, text)
        if n == 2:
            for (t, body) in split_sections(post):
                if t == "Module 11 — Invoicing and Claims":
                    LOG["excluded_blocks"].append({"module": 2, "heading": t, "chars": len(body)})
        mods.append({"num": n, "id": "m%02d" % n, "title": MODULE_TITLES[n], "file": f[0].name,
                     "preamble": preamble, "lessons": lessons, "post": post})

    # course-data.js (structure only — titles and paths; no answer content)
    data = {"modules": [{
        "num": m["num"], "id": m["id"], "title": m["title"], "path": mod_dir(m["num"]) + "/",
        "lessons": [{"id": lesson_id(m["num"], l["lnum"]), "num": "%d.%d" % (m["num"], l["lnum"]),
                     "title": plain(l["title"]), "label": "Lesson %d.%d — %s" % (m["num"], l["lnum"], plain(l["title"])),
                     "path": "%s/%s/" % (mod_dir(m["num"]), lesson_dir(m["num"], l["lnum"]))} for l in m["lessons"]]
    } for m in mods]}
    (OUT / "shared").mkdir(parents=True, exist_ok=True)
    (OUT / "shared" / "course-data.js").write_text(
        "/* NDIS VA Foundations — course structure (generated from the RC4 learner files). Titles and paths only. */\n"
        "window.NDIS_COURSE = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")

    for mi, mod in enumerate(mods):
        d = OUT / mod_dir(mod["num"])
        d.mkdir(parents=True, exist_ok=True)
        (d / "index.html").write_text(build_module(mods, mi), encoding="utf-8")
        for li, les in enumerate(mod["lessons"]):
            ld = d / lesson_dir(mod["num"], les["lnum"])
            ld.mkdir(parents=True, exist_ok=True)
            (ld / "index.html").write_text(build_lesson(mods, mi, li), encoding="utf-8")

    wb = openpyxl.load_workbook(PACK)
    global M15_WB
    M15_WB = openpyxl.load_workbook(sys.argv[3]) if len(sys.argv) > 3 else None
    import pages  # site-level pages (home, progress, wattlebird, certification)
    pages.build_all(OUT, mods, data, wb, sys.modules[__name__])

    report = {
        "modules": [{"num": m["num"], "file": m["file"], "lessons": ["%d.%d %s (%d panels)" % (m["num"], l["lnum"], plain(l["title"]), len(l["steps"])) for l in m["lessons"]],
                     "has_preamble": bool(m["preamble"].strip()), "has_wrapup": bool(m["post"].strip())} for m in mods],
        "log": LOG,
    }
    (Path(__file__).parent / "build_log.json").write_text(json.dumps(report, indent=1, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(LOG, indent=1, ensure_ascii=False))
    print("lessons:", sum(len(m["lessons"]) for m in mods))


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).parent))
    main()
